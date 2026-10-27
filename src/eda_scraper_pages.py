import time
from datetime import datetime

import pandas as pd
from bs4 import BeautifulSoup

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SEARCH_URL = "http://eservices.edaegypt.gov.eg/EDASearch/SearchRegDrugs.aspx"

SEARCH_KEYWORDS = [
    "ABA", "ACE", "ACI", "ACN", "ACT", "ADO", "ADR",
    "ALB", "ALC", "ALE", "ALL", "ALM", "ALP", "ALT",
    "AMO", "AMP", "AMX", "ANA", "AND", "ANT", "API",
    "ARA", "ARI", "ASP", "ATO", "AUG", "AZI",

    "BAC", "BEC", "BEN", "BET", "BIS", "BRO", "BRU",
    "BUD", "BUS",

    "CAL", "CAN", "CAP", "CAR", "CAT", "CEF", "CEH",
    "CEL", "CET", "CIP", "CLA", "CLI", "CLO", "COA",
    "COL", "COR",

    "DEX", "DIA", "DIC", "DIL", "DIM", "DIP", "DOL",
    "DOM", "DOR", "DOX", "DUL",

    "ECO", "EMP", "ENA", "ENO", "ERY", "ESO", "EST",
    "EZE",

    "FAM", "FEN", "FER", "FIN", "FLA", "FLO", "FLU",
    "FOL", "FOR", "FUC",

    "GAB", "GAL", "GEN", "GLI", "GLU",

    "HEP", "HEX", "HIS", "HUM", "HYD", "HYP",

    "IBU", "INS", "IRO", "ISO", "ITR",

    "KET",

    "LAC", "LAM", "LAN", "LAT", "LEV", "LIN", "LIP",
    "LIS", "LOR", "LOS", "LOV",

    "MAC", "MAG", "MEF", "MEL", "MEM", "MER", "MET",
    "MIC", "MIR", "MOM", "MON", "MOT", "MOX", "MUC",

    "NAP", "NAS", "NEB", "NEO", "NIC", "NIF", "NIT",
    "NOR", "NOV", "NYS",

    "OCT", "OLA", "OME", "OND", "OPT", "ORA", "ORS",
    "OXA",

    "PAN", "PAR", "PEN", "PER", "PHA", "POL", "PRA",
    "PRE", "PRO", "PSE", "PUL",

    "QUE", "QUI",

    "RAB", "RAM", "RAN", "REG", "REL", "RIS", "RIV",
    "ROS",

    "SAL", "SER", "SIM", "SIT", "SOL", "SOR", "SPI",
    "STA", "SUP", "SUS",

    "TAM", "TEL", "TER", "THE", "TIA", "TIC", "TOB",
    "TRA", "TRE", "TRI", "TRO",

    "ULC", "URI",

    "VAL", "VAN", "VEN", "VER", "VIL", "VIT", "VOL",

    "WAR",

    "XYL",

    "ZIN", "ZIP", "ZIT", "ZOA", "ZOL", "ZOV"
]

OUTPUT_FILE = "eda_all_drugs.xlsx"
TEMP_FILE = "eda_all_drugs_temp.xlsx"


def clean_text(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split()).strip()


def setup_driver():
    options = Options()
    options.add_argument("--start-maximized")
    driver = webdriver.Chrome(options=options)
    return driver


def extract_basic_fields(soup):
    data = {}

    for group in soup.select(".form-group"):
        label = group.find(["label", "span"])
        input_tag = group.find("input")
        textarea_tag = group.find("textarea")
        select_tag = group.find("select")

        if not label:
            continue

        key = clean_text(label.get_text(" ", strip=True))
        value = ""

        if input_tag:
            value = clean_text(input_tag.get("value", ""))
        elif textarea_tag:
            value = clean_text(textarea_tag.get_text(" ", strip=True))
        elif select_tag:
            selected = select_tag.find("option", selected=True)
            if selected:
                value = clean_text(selected.get_text(" ", strip=True))

        if key and value:
            data[key] = value

    labels = soup.find_all(["label", "span"])

    for label in labels:
        key = clean_text(label.get_text(" ", strip=True))

        if not key or len(key) > 80:
            continue

        parent = label.parent
        if not parent:
            continue

        input_tag = parent.find("input")
        textarea_tag = parent.find("textarea")

        value = ""

        if input_tag:
            value = clean_text(input_tag.get("value", ""))
        elif textarea_tag:
            value = clean_text(textarea_tag.get_text(" ", strip=True))

        if key and value and key not in data:
            data[key] = value

    return data


def table_to_rows(table):
    trs = table.find_all("tr")

    if not trs:
        return []

    headers = []
    first_cells = trs[0].find_all(["th", "td"])

    for i, cell in enumerate(first_cells, start=1):
        header = clean_text(cell.get_text(" ", strip=True))
        headers.append(header if header else f"Column_{i}")

    rows = []

    for tr in trs[1:]:
        cells = tr.find_all(["td", "th"])

        if not cells:
            continue

        row = {}

        for i, cell in enumerate(cells):
            key = headers[i] if i < len(headers) else f"Column_{i + 1}"
            value = clean_text(cell.get_text(" ", strip=True))
            row[key] = value

        if any(row.values()):
            rows.append(row)

    return rows


def extract_tables_as_text(soup):
    data = {}
    tables = soup.find_all("table")

    for index, table in enumerate(tables, start=1):
        table_id = table.get("id", "")
        table_title = table_id if table_id else f"Table_{index}"

        rows = table_to_rows(table)

        if not rows:
            continue

        table_text = []

        for row in rows:
            parts = []

            for key, value in row.items():
                if value:
                    parts.append(f"{key}: {value}")

            if parts:
                table_text.append(" / ".join(parts))

        if table_text:
            data[table_title] = " | ".join(table_text)

    return data


def parse_drug_details_page(html, source_url, search_keyword):
    soup = BeautifulSoup(html, "lxml")

    data = {
        "Search Keyword": search_keyword,
        "Source URL": source_url,
        "Extraction Date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    data.update(extract_basic_fields(soup))
    data.update(extract_tables_as_text(soup))

    return data


def is_real_drug_record(row):
    important_keys = [
        "Name",
        "Product Type",
        "Dosage Form",
        "Registration No",
        "Applicant",
        "Strength",
    ]

    for key in important_keys:
        if clean_text(row.get(key, "")):
            return True

    return False


def find_details_buttons(driver):
    """
    Finds only real drug detail buttons/links.
    It avoids clicking top product category radio buttons.
    """
    buttons = []

    candidates = driver.find_elements(
        By.XPATH,
        "//a | //button | //input[@type='submit'] | //input[@type='button']"
    )

    for el in candidates:
        try:
            input_type = clean_text(el.get_attribute("type")).lower()

            if input_type in ["radio", "checkbox", "text", "hidden"]:
                continue

            text = clean_text(el.text).lower()
            value = clean_text(el.get_attribute("value")).lower()
            title = clean_text(el.get_attribute("title")).lower()
            href = clean_text(el.get_attribute("href")).lower()
            onclick = clean_text(el.get_attribute("onclick")).lower()
            outer = clean_text(el.get_attribute("outerHTML")).lower()
            element_id = clean_text(el.get_attribute("id")).lower()
            name = clean_text(el.get_attribute("name")).lower()

            combined = " ".join([
                text, value, title, href, onclick, outer, element_id, name
            ])

            if text.strip().isdigit():
                continue

            skip_words = [
                "human pharmaceutical",
                "biological",
                "herbal medicine",
                "veterinary",
                "biocides",
                "cosmetics",
                "complementary",
                "captcha",
                "search",
                "clear",
                "reset",
                "download",
                "attachment",
                "facebook",
                "twitter",
                "youtube",
                "instagram",
                "mailto:",
                "tel:",
            ]

            if any(word in combined for word in skip_words):
                continue

            looks_like_details = (
                "detail" in combined
                or "details" in combined
                or "select" in combined
                or "view" in combined
                or "searchregdrugsresult" in combined
                or ("__dopostback" in combined and "grid" in combined)
                or ("__dopostback" in combined and "gv" in combined)
                or ("__dopostback" in combined and "result" in combined)
            )

            if looks_like_details and el.is_displayed() and el.is_enabled():
                buttons.append(el)

        except Exception:
            continue

    unique_buttons = []
    seen = set()

    for btn in buttons:
        try:
            key = btn.get_attribute("outerHTML")
            if key not in seen:
                seen.add(key)
                unique_buttons.append(btn)
        except Exception:
            continue

    return unique_buttons


def collect_current_page_details_by_clicking(driver, search_keyword):
    results = []
    result_page_url = driver.current_url

    detail_count = len(find_details_buttons(driver))
    print(f"Detected details buttons: {detail_count}")

    if detail_count == 0:
        print("No Details buttons found. Checking current page.")

        data = parse_drug_details_page(
            html=driver.page_source,
            source_url=driver.current_url,
            search_keyword=search_keyword,
        )

        if is_real_drug_record(data):
            results.append(data)

        return results

    for index in range(detail_count):
        try:
            print(f"Opening details {index + 1} of {detail_count}")

            buttons = find_details_buttons(driver)

            if index >= len(buttons):
                break

            try:
                buttons[index].click()
            except Exception:
                driver.execute_script("arguments[0].click();", buttons[index])

            time.sleep(3)

            data = parse_drug_details_page(
                html=driver.page_source,
                source_url=driver.current_url,
                search_keyword=search_keyword,
            )

            if is_real_drug_record(data):
                results.append(data)
                print(f"Extracted: {data.get('Name', 'No Name')}")
            else:
                print("Opened page, but no drug details were detected.")

            driver.back()
            time.sleep(3)

        except Exception as e:
            print(f"Error in detail {index + 1}: {e}")

            try:
                driver.get(result_page_url)
                time.sleep(3)
            except Exception:
                pass

    return results


def click_next_page(driver, next_page_number):
    """
    Click numeric pagination pages: 1, 2, 3, 4, 5...
    """
    target = str(next_page_number)

    possible_xpaths = [
        f"//a[normalize-space(text())='{target}']",
        f"//span[normalize-space(text())='{target}']/parent::a",
        f"//input[@value='{target}']",
        f"//button[normalize-space(text())='{target}']",
    ]

    for xpath in possible_xpaths:
        try:
            elements = driver.find_elements(By.XPATH, xpath)

            for el in elements:
                if el.is_displayed() and el.is_enabled():
                    print(f"Moving to results page: {target}")

                    try:
                        el.click()
                    except Exception:
                        driver.execute_script("arguments[0].click();", el)

                    time.sleep(4)
                    return True

        except Exception:
            continue

    print(f"No pagination button found for page: {target}")
    return False


def search_keyword(driver, keyword):
    print("\n" + "=" * 60)
    print(f"Searching for: {keyword}")
    print("=" * 60)

    driver.get(SEARCH_URL)
    time.sleep(3)

    try:
        human_radio = driver.find_element(By.ID, "ContentPlaceHolder1_ui_rdDrugType_0")
        if not human_radio.is_selected():
            human_radio.click()
            time.sleep(1)
    except Exception:
        print("Human Pharmaceutical option was not found. Please select it manually if needed.")

    input_found = False

    trade_name_ids = [
        "ContentPlaceHolder1_ui_txtTradeName",
        "ui_txtTradeName",
    ]

    for input_id in trade_name_ids:
        try:
            trade_input = driver.find_element(By.ID, input_id)
            trade_input.clear()
            trade_input.send_keys(keyword)
            input_found = True
            break
        except Exception:
            continue

    if not input_found:
        print("Trade Name field was not found. Type the keyword manually in Chrome.")

    print("\nManual step in Chrome:")
    print("1. Make sure Human Pharmaceutical is selected.")
    print("2. Enter captcha manually.")
    print("3. Click Search.")
    print("4. Wait until results appear.")
    print("5. Return to PowerShell and press Enter.")

    input("\nPress Enter after results appear...")

    all_results = []
    page_number = 1

    while True:
        print(f"\nExtracting results page: {page_number}")

        page_results = collect_current_page_details_by_clicking(driver, keyword)
        all_results.extend(page_results)

        if all_results:
            df_temp = pd.DataFrame(all_results)
            df_temp = remove_duplicates(df_temp)
            save_professional_excel(df_temp, TEMP_FILE)
            print(f"Temporary file saved: {TEMP_FILE}")

        next_page_number = page_number + 1
        has_next = click_next_page(driver, next_page_number)

        if not has_next:
            break

        page_number += 1

    print(f"Extracted drugs for keyword {keyword}: {len(all_results)}")
    return all_results


def remove_duplicates(df):
    if df.empty:
        return df

    if "Registration No" in df.columns and "Name" in df.columns:
        return df.drop_duplicates(subset=["Registration No", "Name"], keep="first")

    if "Registration No" in df.columns:
        return df.drop_duplicates(subset=["Registration No"], keep="first")

    if "Name" in df.columns:
        return df.drop_duplicates(subset=["Name"], keep="first")

    return df.drop_duplicates()


def save_professional_excel(df, output_file):
    if df.empty:
        return

    preferred_columns = [
        "Name",
        "Product Type",
        "Dosage Form",
        "Shelf Life",
        "Route",
        "Strength",
        "Pack Unit & details",
        "Applicant",
        "Registration No",
        "Registration Type",
        "Marketing Type",
        "Marketing Status",
        "Application Status",
        "License Status",
        "Pricing Status",
        "Registration Data",
        "Combo Pack",
        "Search Keyword",
        "Source URL",
        "Extraction Date",
    ]

    existing_preferred = [col for col in preferred_columns if col in df.columns]
    remaining_columns = [col for col in df.columns if col not in existing_preferred]
    df = df[existing_preferred + remaining_columns]

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Drugs_Data", index=False)

        summary_data = {
            "Metric": [
                "Total Records",
                "Unique Registration Numbers",
                "Unique Drug Names",
                "Unique Applicants",
                "Extraction Date",
            ],
            "Value": [
                len(df),
                df["Registration No"].nunique() if "Registration No" in df.columns else "",
                df["Name"].nunique() if "Name" in df.columns else "",
                df["Applicant"].nunique() if "Applicant" in df.columns else "",
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            ],
        }

        summary_df = pd.DataFrame(summary_data)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

    wb = load_workbook(output_file)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for ws in wb.worksheets:
        ws.sheet_view.rightToLeft = False
        ws.freeze_panes = "A2"

        max_row = ws.max_row
        max_col = ws.max_column

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = thin_border

        ws.auto_filter.ref = ws.dimensions

        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            max_length = 0

            for cell in ws[col_letter]:
                if cell.value:
                    value_length = len(str(cell.value))
                    max_length = max(max_length, value_length)

            adjusted_width = min(max(max_length + 2, 12), 55)
            ws.column_dimensions[col_letter].width = adjusted_width

        ws.row_dimensions[1].height = 30

        for row_idx in range(2, max_row + 1):
            ws.row_dimensions[row_idx].height = 40

        if ws.title == "Drugs_Data" and max_row > 1 and max_col > 1:
            table_ref = f"A1:{get_column_letter(max_col)}{max_row}"
            table = Table(displayName="DrugsDataTable", ref=table_ref)

            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = style

            try:
                ws.add_table(table)
            except Exception:
                pass

    if "Summary" in wb.sheetnames:
        ws = wb["Summary"]
        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 40

    wb.save(output_file)


def main():
    driver = setup_driver()
    all_rows = []

    print(f"Total search keywords: {len(SEARCH_KEYWORDS)}")

    try:
        for index, keyword in enumerate(SEARCH_KEYWORDS, start=1):
            print("\n" + "#" * 70)
            print(f"Starting search {index} of {len(SEARCH_KEYWORDS)}: {keyword}")
            print("#" * 70)

            rows = search_keyword(driver, keyword)
            all_rows.extend(rows)

            if all_rows:
                df_temp = pd.DataFrame(all_rows)
                df_temp = remove_duplicates(df_temp)
                save_professional_excel(df_temp, TEMP_FILE)
                print(f"\nMain temporary file updated: {TEMP_FILE}")

            if index < len(SEARCH_KEYWORDS):
                input("\nPress Enter to continue to next keyword...")

    finally:
        driver.quit()

    if not all_rows:
        print("\nNo data extracted.")
        return

    df = pd.DataFrame(all_rows)
    df = remove_duplicates(df)

    save_professional_excel(df, OUTPUT_FILE)

    print("\nDone.")
    print(f"Final records: {len(df)}")
    print(f"Final Excel file: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()